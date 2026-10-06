"""
rerun_local_test_urdf.py -- fully OFFLINE test.

Loads a URDF + OBJ/MTL export that's already on disk, builds the asset,
creates the joints, simulates, and logs the throw distance. Makes NO
Onshape API calls and has no Onshape plumbing at all: no document-ID
popup, no ID syncing into other scripts, no webhook listener, no webhook
loop. The only network traffic is to the local Kit app on 127.0.0.1
(load asset, create joints, simulate) -- that's how the model gets into
Omniverse, so it can't be removed.

Each run:
    urdf_obj_asset.build_asset()  -> OBJs wrapped in a fresh .usda
    /load_asset + /align_to_ground (+ Z-up stage check)
    urdf_joint_map.run()          -> joint_map.py from the URDF
    create_all_joints.main()      -> joints + dynamic Ball
    collision filtering, arm concave collision, simulate, log distance

To test a different export, change URDF_FOLDER / MESH_FOLDER below.
Run it again to re-test after replacing the files -- NUM_RUNS > 1 repeats
the whole load/simulate cycle back to back in one session.
"""

import os
import glob
import math
import time
import subprocess
import requests
import sys
import openpyxl
from openpyxl.styles import Font
from datetime import datetime

import load_remote
import create_all_joints
import generate_joint_map
import urdf_joint_map
import urdf_obj_asset
import physics_client

# Directory this script lives in (joint_map.py is written here).
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

# -----------------------------------------------------------------------
# PER-RUN FOLDER + BALL DISTANCE LOG -- same convention as
# ExperimentCodeComplete.py: one timestamped folder per script
# execution, one Excel log inside it covering every throw this run
# produces.
# -----------------------------------------------------------------------
#Update with folder location for python scripts
base_loc = r"E:\Onshape test"
current_datetime = datetime.now()
current_datetime_str = current_datetime.strftime("%y_%m_%d_%H_%M")
folder = os.path.join(base_loc, current_datetime_str)
if not os.path.isdir(folder):
    os.makedirs(folder)
BALL_DISTANCE_LOG_PATH = os.path.join(folder, "ball_distance_log.xlsx")


def log_ball_distance(run_label, distance_meters):
    """
    Append one row (run label, distance thrown, timestamp) to
    BALL_DISTANCE_LOG_PATH -- identical logic to
    ExperimentCodeComplete.py's log_ball_distance(), just duplicated
    here since this is a separate script. Creates the file with a
    header row on first call; every call after just appends.
    distance_meters=None logs "N/A" (ball never landed in time).
    """
    if not os.path.isfile(BALL_DISTANCE_LOG_PATH):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Ball Distance Log"
        headers = ["Run", "Distance Thrown (m)", "Timestamp"]
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = Font(name="Arial", bold=True)
        sheet.column_dimensions["A"].width = 16
        sheet.column_dimensions["B"].width = 20
        sheet.column_dimensions["C"].width = 20
        wb.save(BALL_DISTANCE_LOG_PATH)

    wb = openpyxl.load_workbook(BALL_DISTANCE_LOG_PATH)
    sheet = wb["Ball Distance Log"]
    next_row = sheet.max_row + 1
    value = round(distance_meters, 4) if distance_meters is not None else "N/A"
    sheet.cell(row=next_row, column=1, value=run_label).font = Font(name="Arial")
    sheet.cell(row=next_row, column=2, value=value).font = Font(name="Arial")
    sheet.cell(row=next_row, column=3,
               value=datetime.now().strftime("%Y-%m-%d %H:%M:%S")).font = Font(name="Arial")
    wb.save(BALL_DISTANCE_LOG_PATH)
    print(f"Logged to {BALL_DISTANCE_LOG_PATH}: run={run_label}, distance={value}")


# EDIT THESE if your Onshape URDF export lives somewhere else. Folder,
# not filename -- the .urdf inside URDF_FOLDER is found automatically, and
# the URDF's "package://.../meshes/X.obj" paths resolve to MESH_FOLDER/X.obj.
URDF_FOLDER = r"E:\Onshape test\26_09_28_19_48\Catapult_original_urdf\assembly_1\urdf"
MESH_FOLDER = r"E:\Onshape test\26_09_28_19_48\Catapult_original_urdf\assembly_1\meshes"
urdf_joint_map.URDF_FOLDER = URDF_FOLDER
urdf_joint_map.MESH_FOLDER = MESH_FOLDER

# "Z" = Onshape's exact coordinates (Z up) -- same as ExperimentCodeComplete.py.
# Kit's stage must be Z-up (Edit > Preferences > Stage > Default Up Axis = Z,
# then restart Kit); load_and_run() stops with instructions if it isn't.
# "Y" = previous behavior (Onshape Z shown as Kit's Y).
STAGE_UP_AXIS = "Z"
urdf_joint_map.STAGE_UP_AXIS = STAGE_UP_AXIS

# Matches ExperimentCodeComplete.py's own GROUND_PLANE_SIZE -- keep
# these in sync, or the ground plane's real-world size will differ
# between a full pipeline run and this local rerun.
GROUND_PLANE_SIZE = 6000.0

# Per-part mass from the URDF <inertial> blocks (Onshape's own value for
# the assigned material) -- purely local, no API calls. False by default so
# the physics matches the STEP pipeline that worked (PhysX volume x default
# density). Flip to True for real masses. The Ball has no URDF mass either
# way (it's unmated, so the exporter made it the massless root link).
APPLY_URDF_MASS = False

# Frame members (everything jointed EXCEPT these moving parts) get collision
# disabled between each other wherever their bounding boxes touch at all.
# With every frame connection now a pinned/fixed joint, the joints alone
# hold the frame's shape; member-vs-member contact only adds artifacts --
# e.g. VerticalBar's foot sits in a notch of Base with NO mate between the
# two, and Base's convex-decomposition hull fills that notch, so PhysX
# pushed the bar sideways. The parts listed here keep full collision
# against everything (the ball must hit the arm; the arm must be stopped by
# the frame), and every part still collides with the ground.
FILTER_FRAME_SELF_COLLISION = True
MOVING_PART_NAMES = ["ThrowingArm", "Ball"]

# How long (and how often) this script polls /get_ball_landing_result
# after starting the simulation, waiting for the ball to land. Raise
# LANDING_TRACKER_TIMEOUT_S if your throw genuinely takes longer than
# this to land.
LANDING_TRACKER_TIMEOUT_S = 30.0
LANDING_POLL_INTERVAL_S = 0.5


# -----------------------------------------------------------------------
# KIT APP AUTO-LAUNCH
# -----------------------------------------------------------------------
KIT_APP_ROOT = r"C:\OmniverseStream\kit-app-template"

# The GENERATED launcher .bat for this one app -- NOT `.\repo.bat launch`,
# whose interactive app-picker prompt hangs forever with no terminal attached.
KIT_APP_LAUNCHER = os.path.join(
    KIT_APP_ROOT, "_build", "windows-x86_64", "release",
    "my_demo.my_editor_streaming.kit.bat",
)

# Same HTTP server physics_client talks to.
KIT_READY_HOST = physics_client.KIT_HOST
KIT_READY_PORT = physics_client.KIT_PORT

# First boot after a build (cold shader/extension cache) can take minutes.
KIT_BOOT_TIMEOUT_S = 180.0
KIT_READY_POLL_INTERVAL_S = 2.0


def _resolve_kit_app_launcher() -> str:
    """
    Resolve the actual generated launcher .bat under
    _build\\windows-x86_64\\release. Tries the KIT_APP_LAUNCHER guess
    above first; repo.bat build's exact naming can differ slightly
    from the .kit filename (extra suffixes, different casing, etc.),
    so if that guess doesn't exist, this scans the release folder
    itself instead of just failing on a wrong hardcoded name.
    """
    if os.path.isfile(KIT_APP_LAUNCHER):
        return KIT_APP_LAUNCHER

    release_dir = os.path.dirname(KIT_APP_LAUNCHER)
    if not os.path.isdir(release_dir):
        raise FileNotFoundError(
            f"Release folder not found at {release_dir}. Build the app first "
            f"(.\\repo.bat build from {KIT_APP_ROOT})."
        )

    # Filter to actual app launchers: pattern is always "<name>.kit.bat".
    # Excludes noise seen in this project's own release folder --
    # tests-*.bat (test runners, one per app), and one-off utility
    # scripts like kit.bat/pull_kit_sdk.bat/setup_python_env.bat that
    # aren't app launchers at all.
    all_bats = sorted(
        p for p in glob.glob(os.path.join(release_dir, "*.kit.bat"))
        if not os.path.basename(p).lower().startswith("tests-")
    )
    if not all_bats:
        raise FileNotFoundError(
            f"No .bat launcher found in {release_dir} at all. Build the app first "
            f"(.\\repo.bat build from {KIT_APP_ROOT})."
        )

    if len(all_bats) == 1:
        print(f"KIT_APP_LAUNCHER guess ({os.path.basename(KIT_APP_LAUNCHER)}) wasn't found, "
              f"but exactly one .bat exists in {release_dir} -- using it: "
              f"{os.path.basename(all_bats[0])}")
        return all_bats[0]

    raise FileNotFoundError(
        f"Kit app launcher not found at {KIT_APP_LAUNCHER}, and multiple .bat files "
        f"exist in {release_dir}, so this can't guess which one is right:\n  "
        + "\n  ".join(os.path.basename(p) for p in all_bats)
        + f"\n\nUpdate KIT_APP_LAUNCHER above to the correct one from that list."
    )


def launch_kit_app():
    """
    Launch the Kit app in the background and block until its HTTP
    server is actually accepting requests -- the automated equivalent
    of manually running `.\repo.bat launch`, picking the app from its
    prompt, and waiting for the app window to finish loading.

    Uses subprocess.Popen, NOT .run()/check_call(): Kit is a
    long-running server process, not a script that's supposed to exit
    on its own. .run() would block this entire script forever waiting
    for Kit to close.

    If Kit is ALREADY running from a previous session (e.g. left open
    after your last debug run), this just finds its HTTP server
    already responding on the very first poll and returns almost
    immediately -- it never tries to launch a second instance.
    """
    # Skip straight to the ready check if Kit's already up (no point
    # launching a second instance / second port).
    try:
        resp = requests.get(f"http://{KIT_READY_HOST}:{KIT_READY_PORT}/docs", timeout=2)
        if resp.status_code == 200:
            print("Kit app is already running and responding -- skipping launch.")
            return None
    except requests.exceptions.RequestException:
        pass

    launcher_path = _resolve_kit_app_launcher()

    print(f"Launching Kit app: {launcher_path}")
    process = subprocess.Popen(
        [launcher_path],
        cwd=os.path.dirname(launcher_path),
    )

    print(f"Waiting up to {KIT_BOOT_TIMEOUT_S:.0f}s for Kit's HTTP server on "
          f"{KIT_READY_HOST}:{KIT_READY_PORT} to come up...")
    deadline = time.time() + KIT_BOOT_TIMEOUT_S
    while time.time() < deadline:
        if process.poll() is not None:
            raise RuntimeError(
                f"Kit app process exited early (return code {process.returncode}) "
                f"before its HTTP server ever came up -- check the Kit app's own "
                f"console/log window for the actual error (a bad build, a port "
                f"already in use, a missing extension, etc.)."
            )
        try:
            # /docs is FastAPI's auto-generated docs page -- exists the
            # moment uvicorn is listening, regardless of whether a
            # stage/asset has been loaded yet.
            resp = requests.get(f"http://{KIT_READY_HOST}:{KIT_READY_PORT}/docs", timeout=2)
            if resp.status_code == 200:
                print("Kit app is up and responding.")
                return process
        except requests.exceptions.RequestException:
            pass
        time.sleep(KIT_READY_POLL_INTERVAL_S)

    raise TimeoutError(
        f"Kit app didn't respond on {KIT_READY_HOST}:{KIT_READY_PORT} within "
        f"{KIT_BOOT_TIMEOUT_S:.0f}s. It may still be booting (a cold first launch "
        f"after a build/cache-clear can take several minutes) -- check the Kit "
        f"window/log directly, or raise KIT_BOOT_TIMEOUT_S."
    )


def filter_frame_self_collision(entries, run_label):
    """Disable collision for every pair of frame members whose bounding
    boxes touch or overlap at all (threshold 0.0), leaving
    MOVING_PART_NAMES untouched. Pairs that don't touch at rest never
    meet anyway -- the frame is held rigid by its joints."""
    moving = {generate_joint_map._group_key(n) for n in MOVING_PART_NAMES}
    frame = sorted({p for e in entries for p in (e["body0_path"], e["body1_path"])
                    if p and generate_joint_map._group_key(p.rsplit("/", 1)[-1]) not in moving})
    exempt = [(e["body0_path"], e["body1_path"]) for e in entries
              if e["body0_path"] and e["body1_path"]]
    print(f"[{run_label}] Disabling collision between touching frame members "
          f"({len(frame)} parts; keeping full collision for {MOVING_PART_NAMES})...")
    physics_client.auto_filter_overlapping_pairs(frame, exempt, threshold=0.0)


def _ball_world_position(ball_path):
    """(world position in stage units, meters_per_unit) of the Ball prim's
    origin -- its geometric center, see urdf_obj_asset.RECENTER_PART_NAMES."""
    result = physics_client.list_prims(root_path=ball_path)
    for e in result["prims"]:
        if e["path"] == ball_path and e.get("world_translation") is not None:
            return tuple(e["world_translation"]), result.get("meters_per_unit", 1.0)
    return None, result.get("meters_per_unit", 1.0)


def throw_distance_meters(start, end, up_axis, meters_per_unit):
    """Distance ALONG THE GROUND from start to landing (up-axis component
    ignored), always >= 0 regardless of throw direction."""
    up = "XYZ".index(str(up_axis).upper()) if str(up_axis).upper() in ("X", "Y", "Z") else 1
    return math.sqrt(sum((end[a] - start[a]) ** 2 for a in range(3) if a != up)) * meters_per_unit


def load_and_run(urdf_file: str, run_label):
    """
    One full run from the URDF + OBJ export: build the asset, load it,
    joints from the URDF, collision tweaks, optional URDF masses,
    simulate, track and log the ball. Called NUM_RUNS times from main().
    Zero Onshape API calls.

    clear_asset() first -- required on reloads so old joints/bodies and the
    ball tracker's has_landed state don't carry over (see the original
    rerun_local_test.py for the full story).
    """
    physics_client.clear_asset()

    # Fresh filename every run: re-reads the URDF/OBJs (so a refreshed
    # export is picked up) and stops Kit reusing a cached layer.
    safe_label = "".join(c if c.isalnum() else "_" for c in str(run_label))
    usd_file = os.path.join(folder, f"assembly_{safe_label}_{datetime.now().strftime('%H%M%S')}.usda")
    print(f"\n[{run_label}] Building asset from URDF + OBJ: {urdf_file}")
    usd_file, asm = urdf_obj_asset.build_asset(urdf_file, usd_file, MESH_FOLDER)

    print(f"[{run_label}] Loading asset: {usd_file}")
    load_remote.load_asset(usd_file)
    aligned = physics_client.align_to_ground() or {}
    stage_up = str(aligned.get("up_axis", STAGE_UP_AXIS)).upper()
    if stage_up != STAGE_UP_AXIS:
        raise RuntimeError(
            f"Kit's stage is {stage_up}-up, but the asset was built for a {STAGE_UP_AXIS}-up stage "
            f"(STAGE_UP_AXIS = {STAGE_UP_AXIS!r}) -- the model would be lying on its side. "
            f"Either set Kit's stage to {STAGE_UP_AXIS}-up (Edit > Preferences > Stage > "
            f"Default Up Axis, then restart Kit) or set STAGE_UP_AXIS = {stage_up!r} here.")

    print(f"[{run_label}] Generating joint_map.py from the URDF...")
    urdf_joint_map.run(write=True, urdf_path=urdf_file)

    import joint_map
    print(f"[{run_label}] Creating joints from joint_map.py...")
    create_all_joints.main()

    print(f"[{run_label}] Filtering overlapping-but-unjointed collision pairs (no Onshape calls)...")
    generate_joint_map.disable_mechanism_self_collision(joint_map.JOINT_MAP)
    if FILTER_FRAME_SELF_COLLISION:
        filter_frame_self_collision(joint_map.JOINT_MAP, run_label)

    throwing_arm_path = generate_joint_map.find_part_prim_path_offline("ThrowingArm")
    if throwing_arm_path:
        # convexDecomposition -- tried "sdf" instead, but that's not a
        # valid token in this Kit/USD schema version ("Unknown
        # approximation token: 'sdf'"). Back to convexDecomposition
        # with the existing voxelResolution/maxConvexHulls tuning below.
        physics_client.set_collision_approximation(throwing_arm_path, "convexDecomposition")
    else:
        print(f"[{run_label}] WARNING: could not find 'ThrowingArm' among live prims -- "
              f"skipping its concave-collision override.")

    if APPLY_URDF_MASS:
        for e in physics_client.list_prims(root_path="/World/ImportedAsset")["prims"]:
            leaf = e["path"].rsplit("/", 1)[-1]
            link = next((n for n in asm.parts if asm.instance_name[n] == leaf), None)
            mass = asm.links[link]["mass"] if link else None
            if mass and mass > 1e-6:
                physics_client.set_mass(e["path"], mass)
    else:
        print(f"[{run_label}] APPLY_URDF_MASS=False -- PhysX default volume x density masses.")

    ball_path = generate_joint_map.find_part_prim_path_offline("Ball")
    start_pos, meters_per_unit, up_axis = None, 1.0, STAGE_UP_AXIS
    if ball_path:
        tracker = physics_client.start_ball_landing_tracker(ball_path) or {}
        up_axis = tracker.get("up_axis", STAGE_UP_AXIS)
        start_pos, meters_per_unit = _ball_world_position(ball_path)  # before the throw
    else:
        print(f"[{run_label}] WARNING: could not resolve a 'Ball' part -- "
              f"landing won't be tracked.")

    physics_client.start_simulation()

    if ball_path:
        print(f"[{run_label}] Waiting up to {LANDING_TRACKER_TIMEOUT_S:.0f}s for the ball "
              f"to land (polling every {LANDING_POLL_INTERVAL_S:.1f}s)...")
        deadline = time.time() + LANDING_TRACKER_TIMEOUT_S
        landed = False
        while time.time() < deadline:
            result = physics_client.get_ball_landing_result(ball_path)
            if result.get("has_landed"):
                landed = True
                # Kit pauses the sim at landing, so the current position IS the landing spot.
                end_pos, _ = _ball_world_position(ball_path)
                distance = None
                if start_pos and end_pos and math.dist(start_pos, end_pos) > 1e-6:
                    distance = throw_distance_meters(start_pos, end_pos, up_axis, meters_per_unit)
                    print(f"[{run_label}] Ball started at {tuple(round(v, 3) for v in start_pos)} "
                          f"and landed at {tuple(round(v, 3) for v in end_pos)} (stage units).")
                elif start_pos:
                    distance = abs(result["landing_x"] - start_pos[0]) * meters_per_unit
                    print(f"[{run_label}] (Ball's stage position didn't update -- using Kit's landing X only.)")
                if distance is None:
                    distance = abs(result["distance_meters"])
                print(f"[{run_label}] Distance thrown: {distance:.4f} m  "
                      f"(Kit's own signed value was {result['distance_meters']:.4f} m)")
                log_ball_distance(run_label, distance)
                break
            time.sleep(LANDING_POLL_INTERVAL_S)
        if not landed:
            print(f"[{run_label}] Ball hadn't landed after {LANDING_TRACKER_TIMEOUT_S:.0f}s "
                  f"of polling -- either the throw takes longer than that, the mechanism "
                  f"didn't launch it, or it's not crossing back down to the ground plane's "
                  f"height at all.")
            log_ball_distance(run_label, None)


# How many back-to-back load/simulate/log cycles to run in one session
# (each one clears the stage first). 1 = a single test run.
NUM_RUNS = 1


def main():
    urdf_file = urdf_joint_map.find_urdf_file(URDF_FOLDER)
    print(f"Using URDF: {urdf_file}\nMeshes: {MESH_FOLDER}")

    launch_kit_app()  # reuses Kit if it's already running

    load_remote.create_ground_plane(size=GROUND_PLANE_SIZE)

    for run in range(1, NUM_RUNS + 1):
        load_and_run(urdf_file, run_label="original" if run == 1 else f"repeat #{run - 1}")

    print(f"\nDone -- {NUM_RUNS} run(s). Results logged to {BALL_DISTANCE_LOG_PATH}")


if __name__ == "__main__":
    main()
